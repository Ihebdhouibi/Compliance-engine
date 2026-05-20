import sys, json, io
from openpyxl import load_workbook

out = open(sys.argv[1], 'w', encoding='utf-8')
def p(*a):
    out.write(" ".join(str(x) for x in a) + "\n")

for path in sys.argv[2:]:
    p("=" * 100)
    p("FILE:", path)
    wb = load_workbook(path, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        p("-" * 100)
        p(f"SHEET: {sn} (rows={ws.max_row}, cols={ws.max_column})")
        for r in ws.iter_rows(values_only=True):
            # Skip fully empty rows
            if not any(c not in (None, "") for c in r):
                continue
            p(list(r))
out.close()

