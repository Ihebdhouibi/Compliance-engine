"""
OCR engine with unified table extraction using pdfplumber + RapidLayout+RapidTable.
Supports images, PDFs, DOCX (native tables + embedded images), and XLSX.
"""
from __future__ import annotations

import io
import logging
import os
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, List

import fitz  # PyMuPDF
import numpy as np
import pandas as pd
from PIL import Image

logger = logging.getLogger(__name__)

# ── RapidOCR (plain text) ──────────────────────────────────────────────
_OCR = None

def _get_ocr():
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR()
    return _OCR

# ── Table extraction engines (RapidLayout + RapidTable) ──────────────
_LAYOUT_ENGINE = None
_TABLE_ENGINE = None

def _get_layout_engine():
    global _LAYOUT_ENGINE
    if _LAYOUT_ENGINE is None:
        try:
            from rapid_layout import RapidLayout, RapidLayoutInput
            from rapid_layout.utils.typings import ModelType as LayoutModelType
            cfg = RapidLayoutInput(model_type=LayoutModelType.PP_LAYOUT_TABLE)
            _LAYOUT_ENGINE = RapidLayout(cfg)
            logger.info("RapidLayout initialized (table detection)")
        except ImportError:
            logger.warning("rapid-layout not installed. Table detection will be limited.")
        except Exception as e:
            logger.warning(f"RapidLayout init failed: {e}")
    return _LAYOUT_ENGINE

def _get_table_engine():
    global _TABLE_ENGINE
    if _TABLE_ENGINE is None:
        try:
            from rapid_table import RapidTable, RapidTableInput, ModelType as TableModelType
            cfg = RapidTableInput(model_type=TableModelType.SLANETPLUS, use_ocr=True)
            _TABLE_ENGINE = RapidTable(cfg)
            logger.info("RapidTable initialized (table recognition)")
        except ImportError:
            logger.warning("rapid-table not installed. Table extraction unavailable.")
        except Exception as e:
            logger.warning(f"RapidTable init failed: {e}")
    return _TABLE_ENGINE

# ── Helper functions for table extraction ──────────────────────────────

def _crop_image(image_array: np.ndarray, box: List[float], pad: int = 4) -> np.ndarray:
    h, w = image_array.shape[:2]
    x1, y1, x2, y2 = box
    x1 = max(int(x1) - pad, 0)
    y1 = max(int(y1) - pad, 0)
    x2 = min(int(x2) + pad, w)
    y2 = min(int(y2) + pad, h)
    return image_array[y1:y2, x1:x2]

def _html_to_dataframe(html: str) -> Optional[pd.DataFrame]:
    if not html:
        return None
    try:
        parsed = pd.read_html(io.StringIO(html))
        return parsed[0] if parsed else None
    except ValueError:
        return None

def extract_tables_from_image(image: Image.Image | np.ndarray) -> List[pd.DataFrame]:
    """Find and recognize every table in a single image."""
    if isinstance(image, Image.Image):
        image_array = np.array(image.convert("RGB"))
    else:
        image_array = image

    layout = _get_layout_engine()
    table_engine = _get_table_engine()
    if layout is None or table_engine is None:
        return []

    try:
        layout_result = layout(image_array)
        boxes = layout_result.boxes or []
        class_names = layout_result.class_names or []

        table_crops = [
            _crop_image(image_array, box)
            for box, cls in zip(boxes, class_names)
            if "table" in cls.lower()
        ]
        if not table_crops:
            table_crops = [image_array]

        dfs = []
        for crop in table_crops:
            table_out = table_engine(crop)
            if table_out is None:
                continue
            for html in table_out.pred_htmls or []:
                df = _html_to_dataframe(html)
                if df is not None:
                    dfs.append(df)
        return dfs
    except Exception as e:
        logger.warning(f"Table extraction from image failed: {e}")
        return []

def extract_tables_from_pdf_images(pdf_path: str | Path, dpi: int = 200) -> List[pd.DataFrame]:
    """Render PDF pages to images and extract tables using RapidLayout+RapidTable."""
    pdf_path = Path(pdf_path)
    try:
        from pdf2image import convert_from_path
        page_images = convert_from_path(str(pdf_path), dpi=dpi)
    except Exception:
        # fallback to PyMuPDF
        import fitz
        doc = fitz.open(str(pdf_path))
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)
        page_images = []
        for page in doc:
            pix = page.get_pixmap(matrix=matrix)
            page_images.append(Image.open(io.BytesIO(pix.tobytes("png"))))
        doc.close()

    all_tables = []
    for page_img in page_images:
        all_tables.extend(extract_tables_from_image(page_img))
    return all_tables

def _dedupe_row_cells(row) -> List[str]:
    """Collapse a python-docx table row into its true (unmerged) cell values.

    python-docx repeats the *same* underlying cell for every column a
    horizontal merge spans (e.g. a "DOMAIN 1" header merged across 8
    columns comes back as that text 8 times). We detect that by identity
    of the underlying XML element (`cell._tc`) and collapse consecutive
    duplicates into a single value, so a merged header becomes one cell
    instead of N repeats.
    """
    values: List[str] = []
    prev_tc = None
    for cell in row.cells:
        if prev_tc is not None and cell._tc is prev_tc:
            continue  # same merged cell as the previous column, skip
        values.append(cell.text.strip())
        prev_tc = cell._tc
    return values


def _docx_table_to_dataframe(table: "DocxTable") -> pd.DataFrame:
    """Convert a python-docx Table into a clean, readable DataFrame.

    - Merged cells are collapsed once (see `_dedupe_row_cells`) instead of
      being repeated across every spanned column.
    - A row that collapses to a single value (i.e. the *entire* row was
      one merged cell — typically a section/domain header inside a big
      table) is kept as its own labeled row instead of being padded with
      duplicate text, so section breaks stay readable rather than noisy.
    - All rows are padded/truncated to the header's column count so the
      DataFrame stays rectangular even when merge patterns differ row to row.
    """
    raw_rows = [_dedupe_row_cells(row) for row in table.rows]
    raw_rows = [r for r in raw_rows if any(v for v in r)]  # drop fully-empty rows
    if not raw_rows:
        return pd.DataFrame()

    header, *body = raw_rows
    n_cols = len(header) or 1

    def _pad(row: List[str]) -> List[str]:
        if len(row) == 1 and n_cols > 1:
            # Whole row was one merged cell (a section/domain header row) —
            # keep the label in the first column, leave the rest blank.
            return [row[0]] + [""] * (n_cols - 1)
        if len(row) < n_cols:
            return row + [""] * (n_cols - len(row))
        return row[:n_cols]

    body = [_pad(r) for r in body]
    df = pd.DataFrame(body, columns=header if header else None)
    return df


def extract_tables_from_docx(docx_path: str | Path, include_embedded_images: bool = True) -> List[pd.DataFrame]:
    """Extract tables from DOCX: native tables + embedded images."""
    docx_path = Path(docx_path)
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed. DOCX table extraction unavailable.")
        return []

    document = Document(str(docx_path))
    tables = []

    # Native tables
    for table in document.tables:
        df = _docx_table_to_dataframe(table)
        if not df.empty:
            tables.append(df)

    # Embedded images
    if include_embedded_images:
        for rel in document.part.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                blob = rel.target_part.blob
                img = Image.open(io.BytesIO(blob)).convert("RGB")
                tables.extend(extract_tables_from_image(img))
            except Exception as e:
                logger.warning(f"Failed to process embedded image: {e}")

    return tables

# ── Data models ─────────────────────────────────────────────────────────

@dataclass
class OcrPage:
    page: int
    text: str
    confidence: float = 0.0

@dataclass
class OcrResult:
    pages: list[OcrPage] = field(default_factory=list)
    engine: str = "rapidocr"
    elapsed_ms: int = 0

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def page_count(self) -> int:
        return len(self.pages)

# ── Main OCR Engine ────────────────────────────────────────────────────

class OcrEngine:
    def __init__(self) -> None:
        self._initialized = False
        self._temp_dir = os.environ.get("OCR_TEMP_DIR") or os.path.join(
            tempfile.gettempdir(), "compliance-engine-ocr"
        )
        self._pdf_dpi = int(os.environ.get("OCR_PDF_DPI", "200"))
        self._enable_table = os.environ.get("ENABLE_TABLE_EXTRACTION", "true").lower() in ("true", "1", "yes")

    def initialize(self) -> None:
        if self._initialized:
            return
        os.makedirs(self._temp_dir, exist_ok=True)
        try:
            _get_ocr()
        except Exception as e:
            logger.warning(f"RapidOCR warmup failed: {e}")
        if self._enable_table:
            _get_layout_engine()
            _get_table_engine()
        self._initialized = True
        logger.info("OCR engine initialized (table extraction %s)",
                    "enabled" if self._enable_table else "disabled")

    def recognize(self, file_path: str, mime: Optional[str] = None) -> OcrResult:
        if not self._initialized:
            self.initialize()
        if not os.path.exists(file_path):
            raise FileNotFoundError(file_path)

        start = time.time()
        kind = self._kind(file_path, mime)
        logger.info(f"OCR start file={file_path} kind={kind}")

        if kind == "image":
            pages = self._ocr_image(file_path)
        elif kind == "pdf":
            pages = self._ocr_pdf(file_path)
        elif kind == "docx":
            pages = self._ocr_docx(file_path)
        elif kind == "xlsx":
            pages = self._ocr_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_path} (mime={mime})")

        elapsed = int((time.time() - start) * 1000)
        return OcrResult(pages=pages, engine="rapidocr", elapsed_ms=elapsed)

    def _kind(self, file_path: str, mime: Optional[str]) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        if ext in (".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"):
            return "image"
        if ext == ".pdf":
            return "pdf"
        if ext == ".docx":
            return "docx"
        if ext in (".xlsx", ".xls"):
            return "xlsx"
        if mime:
            m = mime.lower()
            if m.startswith("image/"):
                return "image"
            if m == "application/pdf":
                return "pdf"
            if "spreadsheetml" in m or "ms-excel" in m:
                return "xlsx"
            if "wordprocessingml" in m or "officedocument" in m:
                return "docx"
        return "unknown"

    # ── Image OCR ─────────────────────────────────────────────────────────

    def _ocr_image(self, file_path: str) -> list[OcrPage]:
        result, _ = _get_ocr()(file_path)
        text = self._flatten_rapid_result(result)

        if self._enable_table:
            try:
                img = Image.open(file_path)
                dfs = extract_tables_from_image(img)
                if dfs:
                    section = self._render_table_group("Tables (detected via layout + structure recognition)", dfs)
                    text = (text + "\n\n" + section).strip()
            except Exception as e:
                logger.warning(f"Image table extraction failed: {e}")

        return [OcrPage(page=1, text=text, confidence=0.0)]

    # ── PDF OCR ──────────────────────────────────────────────────────────

    def _ocr_pdf(self, file_path: str) -> list[OcrPage]:
        pages = []
        doc = fitz.open(file_path)

        # 1. Try pdfplumber for text-based tables (fast, accurate)
        pdfplumber_tables = self._extract_tables_from_pdf_plumber(file_path)
        if pdfplumber_tables:
            table_section = "\n\n" + self._render_table_group(
                "Tables (extracted from PDF text layer)", pdfplumber_tables
            )
        else:
            table_section = ""

        # 2. If no tables found, fallback to RapidLayout on rendered images
        if not pdfplumber_tables and self._enable_table:
            try:
                image_tables = extract_tables_from_pdf_images(file_path, dpi=self._pdf_dpi)
                if image_tables:
                    table_section = "\n\n" + self._render_table_group(
                        "Tables (detected via layout + structure recognition)", image_tables
                    )
            except Exception as e:
                logger.warning(f"PDF table extraction with RapidLayout failed: {e}")

        # For each page, get plain text (native or OCR)
        for i, pdf_page in enumerate(doc, start=1):
            text = pdf_page.get_text().strip()
            if text:
                page_text = text
            else:
                # Fallback to image OCR
                zoom = self._pdf_dpi / 72.0
                mat = fitz.Matrix(zoom, zoom)
                pix = pdf_page.get_pixmap(matrix=mat, alpha=False)
                tmp_img = os.path.join(self._temp_dir, f"pdfpage_{uuid.uuid4().hex}.png")
                pix.save(tmp_img)
                try:
                    result, _ = _get_ocr()(tmp_img)
                    page_text = self._flatten_rapid_result(result)
                finally:
                    try:
                        os.remove(tmp_img)
                    except OSError:
                        pass

            # Append the table section to each page (it contains all tables)
            page_text = page_text + table_section

            pages.append(OcrPage(page=i, text=page_text, confidence=1.0 if text else 0.0))

        doc.close()
        return pages

    def _extract_tables_from_pdf_plumber(self, file_path: str) -> List[pd.DataFrame]:
        """Extract tables from a PDF using pdfplumber, returning a list of DataFrames."""
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber not installed. Text-based PDF table extraction unavailable.")
            return []

        try:
            with pdfplumber.open(file_path) as pdf:
                all_tables = []
                for page in pdf.pages:
                    tables = page.extract_tables()
                    for table in tables:
                        if not table or len(table) < 2:
                            continue
                        # Convert to DataFrame
                        rows = []
                        for row in table:
                            cells = [str(c) if c is not None else "" for c in row]
                            rows.append(cells)
                        if rows:
                            header, *body = rows
                            df = pd.DataFrame(body, columns=header if header else None)
                            if not df.empty:
                                all_tables.append(df)
                return all_tables
        except Exception as e:
            logger.warning(f"pdfplumber table extraction failed: {e}")
            return []

    # ── DOCX OCR ─────────────────────────────────────────────────────────

    def _ocr_docx(self, file_path: str) -> list[OcrPage]:
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx not installed. DOCX processing unavailable.")
            return [OcrPage(page=1, text="", confidence=0.0)]

        doc = Document(file_path)

        # 1. Narrative text: paragraphs only. Table content is handled
        #    separately below (as clean, deduplicated tables) instead of
        #    being flattened into pipe-joined lines here — flattening
        #    merged-cell tables that way just repeats the same header
        #    text once per spanned column and reads as noise.
        paragraph_lines = [
            t for p in doc.paragraphs if (t := (p.text or "").strip())
        ]
        text = "\n".join(paragraph_lines)

        # 2. Tables: native Word tables (merge-aware, deduplicated) and,
        #    optionally, tables recovered from pasted-in images — kept as
        #    two clearly labeled, non-overlapping groups so it's obvious
        #    which extraction path produced which table.
        if self._enable_table:
            native_tables = [
                df for t in doc.tables
                if not (df := _docx_table_to_dataframe(t)).empty
            ]

            image_tables: List[pd.DataFrame] = []
            for rel in doc.part.rels.values():
                if "image" not in rel.reltype:
                    continue
                try:
                    blob = rel.target_part.blob
                    img = Image.open(io.BytesIO(blob)).convert("RGB")
                    image_tables.extend(extract_tables_from_image(img))
                except Exception as e:
                    logger.warning(f"Failed to process embedded image: {e}")

            table_sections = []
            if native_tables:
                table_sections.append(
                    self._render_table_group("Word Tables (native)", native_tables)
                )
            if image_tables:
                table_sections.append(
                    self._render_table_group("Tables Recovered from Embedded Images", image_tables)
                )

            if table_sections:
                text = (text + "\n\n" + "\n\n".join(table_sections)).strip()

        if text:
            return [OcrPage(page=1, text=text, confidence=1.0)]
        else:
            return [OcrPage(page=1, text="", confidence=0.0)]

    @staticmethod
    def _render_table_group(label: str, dfs: List[pd.DataFrame]) -> str:
        """Render a group of tables as one clearly-labeled Markdown section,
        with each table numbered so multiple tables under the same label
        (e.g. several native Word tables) stay distinguishable.
        """
        blocks = [f"## {label}"]
        for i, df in enumerate(dfs, start=1):
            heading = f"### Table {i}" if len(dfs) > 1 else "### Table"
            try:
                body = df.to_markdown(index=False)
            except ImportError:
                body = df.to_string(index=False)
            blocks.append(f"{heading}\n\n{body}")
        return "\n\n".join(blocks)

    # ── XLSX OCR ─────────────────────────────────────────────────────────

    def _ocr_xlsx(self, file_path: str) -> list[OcrPage]:
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not installed. XLSX extraction unavailable.")
            return [OcrPage(page=1, text="", confidence=0.0)]

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            logger.warning(f"XLSX open failed: {e}")
            return [OcrPage(page=1, text="", confidence=0.0)]

        pages = []
        try:
            for idx, ws in enumerate(wb.worksheets, start=1):
                rows_text = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        rows_text.append(" | ".join(cells))
                sheet_text = "\n".join(rows_text).strip()
                if sheet_text:
                    pages.append(OcrPage(page=idx, text=f"[Sheet: {ws.title}]\n{sheet_text}", confidence=1.0))
        finally:
            wb.close()

        if not pages:
            pages.append(OcrPage(page=1, text="", confidence=0.0))
        return pages

    @staticmethod
    def _flatten_rapid_result(result) -> str:
        if not result:
            return ""
        texts = []
        for line in result:
            if not line or len(line) < 2:
                continue
            text_info = line[1]
            if isinstance(text_info, (tuple, list)) and len(text_info) >= 1:
                text = str(text_info[0])
                if text:
                    texts.append(text)
        return "\n".join(texts)

# ── Singleton ──────────────────────────────────────────────────────────

_engine: OcrEngine | None = None

def get_engine() -> OcrEngine:
    global _engine
    if _engine is None:
        _engine = OcrEngine()
    return _engine
